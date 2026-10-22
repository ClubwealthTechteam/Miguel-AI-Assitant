# Usage: python3 make_pptx.py <images_dir> <out.pptx> "<title>" "<subtitle>"
import sys, glob, os
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from PIL import Image
src, out, title, sub = sys.argv[1:5]
NAMES = {'0-previous-day-event': "Previous day's event", '1-upcoming': 'This week', '2-next-tour': 'Next Half-Day tour',
         '3-daily-sign-ups': 'Daily sign-ups', '4-who-brought-them-in': 'Who brought them in',
         '5-campaign-codes': 'Campaign codes', '6-cwsocial27': 'CWSocial27', '2-upcoming-next-tour': 'Next tour'}
p = Presentation(); p.slide_width, p.slide_height = Inches(13.333), Inches(7.5)
W, H = p.slide_width, p.slide_height
s = p.slides.add_slide(p.slide_layouts[6])
for txt, top, size, bold, col in [(title, 2.6, 40, True, '1F2937'), (sub, 3.6, 20, False, '5B6472')]:
    tb = s.shapes.add_textbox(Inches(0.8), Inches(top), W - Inches(1.6), Inches(1)); r = tb.text_frame.paragraphs[0].add_run()
    r.text, r.font.size, r.font.bold, r.font.color.rgb = txt, Pt(size), bold, RGBColor.from_string(col)
for f in sorted(glob.glob(os.path.join(src, '*.png'))):
    key = os.path.basename(f)[:-4]
    s = p.slides.add_slide(p.slide_layouts[6])
    tb = s.shapes.add_textbox(Inches(0.5), Inches(0.25), W - Inches(1), Inches(0.6)); r = tb.text_frame.paragraphs[0].add_run()
    r.text, r.font.size, r.font.bold, r.font.color.rgb = NAMES.get(key, key), Pt(24), True, RGBColor.from_string('1F2937')
    iw, ih = Image.open(f).size
    maxw, maxh = W - Inches(1), H - Inches(1.3)
    sc = min(maxw / iw, maxh / ih); w, h = int(iw * sc), int(ih * sc)
    s.shapes.add_picture(f, int((W - w) / 2), Inches(1.0) + int((maxh - h) / 2), w, h)
p.save(out); print(out, len(p.slides))
